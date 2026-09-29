import io
import json
import sqlite3
import tempfile
import unittest
import zipfile
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch
from urllib.parse import urlsplit

from openpyxl import load_workbook

from app import (
    Household,
    HouseholdInvitation,
    OdometerEntry,
    Reminder,
    ServiceRecord,
    User,
    ValueEntry,
    Vehicle,
    create_app,
    db,
    reminder_status,
    valid_discord_webhook_url,
)
from werkzeug.security import generate_password_hash


class CarFreeksTestCase(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        data_dir = Path(self.temp_dir.name)
        self.app = create_app(
            {
                "TESTING": True,
                "SECRET_KEY": "test-secret-key-that-is-long-enough",
                "CARFREEKS_DATA_DIR": str(data_dir),
                "UPLOAD_FOLDER": str(data_dir / "receipts"),
                "SQLALCHEMY_DATABASE_URI": f"sqlite:///{(data_dir / 'test.db').as_posix()}",
                "WTF_CSRF_ENABLED": False,
            }
        )
        self.client = self.app.test_client()

    def tearDown(self):
        with self.app.app_context():
            db.session.remove()
            db.engine.dispose()
        self.temp_dir.cleanup()

    def login(self):
        response = self.client.post(
            "/register",
            data={
                "username": "owner",
                "password": "test-password-long",
                "confirm_password": "test-password-long",
            },
            follow_redirects=True,
        )
        if (
            response.status_code != 200
            or b"Your CarFreeks account is ready." not in response.data
        ):
            response = self.client.post(
                "/login",
                data={"username": "owner", "password": "test-password-long"},
                follow_redirects=True,
            )
        return response

    def test_signing_key_is_generated_and_persisted_without_environment_setup(self):
        data_dir = Path(self.temp_dir.name) / "automatic-key"
        config = {
            "TESTING": True,
            "CARFREEKS_DATA_DIR": str(data_dir),
            "SQLALCHEMY_DATABASE_URI": f"sqlite:///{(data_dir / 'test.db').as_posix()}",
            "WTF_CSRF_ENABLED": False,
        }
        with patch.dict("os.environ", {"CARFREEKS_SECRET_KEY": "legacy-stack-setting"}):
            first_app = create_app(config)
            first_key = first_app.config["SECRET_KEY"]
            self.assertGreaterEqual(len(first_key), 40)
            self.assertNotEqual(first_key, "legacy-stack-setting")
            self.assertEqual(
                (data_dir / "flask-secret-key").read_text(encoding="utf-8"),
                first_key,
            )
            with first_app.app_context():
                db.session.remove()
                db.engine.dispose()

            second_app = create_app(config)
            self.assertEqual(second_app.config["SECRET_KEY"], first_key)
            with second_app.app_context():
                db.session.remove()
                db.engine.dispose()

    def test_first_run_owner_account_setup(self):
        response = self.client.get("/login")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Create an account", response.data)

        setup = self.client.get("/register")
        self.assertEqual(setup.status_code, 200)
        self.assertIn(b"Create your account", setup.data)

        mismatch = self.client.post(
            "/register",
            data={
                "username": "owner",
                "password": "test-password-long",
                "confirm_password": "different-password",
            },
        )
        self.assertEqual(mismatch.status_code, 200)
        self.assertIn(b"Passwords must match", mismatch.data)

        created = self.client.post(
            "/register",
            data={
                "username": "owner",
                "password": "test-password-long",
                "confirm_password": "test-password-long",
            },
        )
        self.assertEqual(created.status_code, 302)
        self.assertEqual(created.headers["Location"], "/")
        self.assertEqual(self.client.get("/").status_code, 200)
        self.assertEqual(self.add_vehicle().status_code, 302)

        self.client.post("/logout")
        login_page = self.client.get("/login")
        self.assertEqual(login_page.status_code, 200)
        self.assertIn(b"Create an account", login_page.data)
        self.assertEqual(self.client.get("/register").status_code, 200)

        second_vehicle_id = None
        with self.app.test_client() as another_client:
            duplicate_username = another_client.post(
                "/register",
                data={
                    "username": "owner",
                    "password": "another-password-long",
                    "confirm_password": "another-password-long",
                },
                follow_redirects=True,
            )
            self.assertEqual(duplicate_username.status_code, 200)
            self.assertIn(b"That username is already in use.", duplicate_username.data)

            another_client.post(
                "/register",
                data={
                    "username": "second-owner",
                    "password": "another-password-long",
                    "confirm_password": "another-password-long",
                },
            )
            second_client_vehicle = another_client.post(
                "/vehicles/new",
                data={
                    "nickname": "Private coupe",
                    "vehicle_type": "Car",
                    "year": "",
                    "make": "",
                    "model": "",
                    "vin": "",
                    "license_plate": "",
                },
            )
            self.assertEqual(second_client_vehicle.status_code, 302)
            with self.app.app_context():
                second_vehicle_id = Vehicle.query.filter_by(nickname="Private coupe").one().id

        self.login()
        dashboard = self.client.get("/")
        self.assertIn(b"Bluebird", dashboard.data)
        self.assertNotIn(b"Private coupe", dashboard.data)
        self.assertEqual(self.client.get(f"/vehicles/{second_vehicle_id}").status_code, 404)
        with self.app.app_context():
            self.assertEqual(User.query.count(), 2)

    def test_owner_can_change_username_and_password(self):
        self.login()

        page = self.client.get("/account")
        self.assertEqual(page.status_code, 200)
        self.assertIn(b"Update your account", page.data)

        incorrect = self.client.post(
            "/account",
            data={
                "current_password": "incorrect-current-password",
                "username": "new-owner",
                "password": "new-password-long",
                "confirm_password": "new-password-long",
            },
        )
        self.assertEqual(incorrect.status_code, 200)
        self.assertIn(b"current password is incorrect", incorrect.data)

        updated = self.client.post(
            "/account",
            data={
                "current_password": "test-password-long",
                "username": "new-owner",
                "password": "new-password-long",
                "confirm_password": "new-password-long",
            },
        )
        self.assertEqual(updated.status_code, 302)
        self.assertEqual(updated.headers["Location"], "/account")
        self.client.post("/logout")

        login = self.client.post(
            "/login",
            data={"username": "new-owner", "password": "new-password-long"},
            follow_redirects=True,
        )
        self.assertIn(b"Your garage.", login.data)

    def test_household_invite_grants_shared_private_garage_once(self):
        self.login()
        self.add_vehicle()
        created = self.client.post("/household/invite")
        self.assertEqual(created.status_code, 302)
        household_page = self.client.get("/household")
        self.assertEqual(household_page.status_code, 200)
        self.assertIn(b"Share this invite link", household_page.data)
        invite_url = urlsplit(
            self._extract_invite_url(household_page.data.decode("utf-8"))
        )

        with self.app.test_client() as invited_client:
            registration = invited_client.get(invite_url.path + "?" + invite_url.query)
            self.assertIn(b"You are joining owner&#39;s Garage", registration.data)
            joined = invited_client.post(
                invite_url.path + "?" + invite_url.query,
                data={
                    "username": "family-member",
                    "password": "family-password-long",
                    "confirm_password": "family-password-long",
                    "invite_token": invite_url.query.removeprefix("invite="),
                },
                follow_redirects=True,
            )
            self.assertIn(b"Bluebird", joined.data)
            self.assertEqual(invited_client.get("/exports/backup.zip").status_code, 403)

            invited_client.post("/logout")
            replay = invited_client.post(
                invite_url.path + "?" + invite_url.query,
                data={
                    "username": "third-person",
                    "password": "third-password-long",
                    "confirm_password": "third-password-long",
                    "invite_token": invite_url.query.removeprefix("invite="),
                },
                follow_redirects=True,
            )
            self.assertIn(b"invalid or has expired", replay.data)

            invited_client.post(
                "/login",
                data={
                    "username": "family-member",
                    "password": "family-password-long",
                },
            )
            self.assertEqual(invited_client.post("/household/invite").status_code, 403)
            vehicle_id = Vehicle.query.one().id
            self.assertEqual(
                invited_client.get(f"/vehicles/{vehicle_id}").status_code, 200
            )

        with self.app.app_context():
            self.assertEqual(User.query.count(), 2)
            self.assertEqual(HouseholdInvitation.query.count(), 1)

    @staticmethod
    def _extract_invite_url(page):
        marker = 'id="invite-link" type="text" readonly value="'
        return page.split(marker, 1)[1].split('"', 1)[0]

    def test_theme_can_be_toggled_and_is_per_user(self):
        self.login()
        theme = self.client.post("/theme", data={"theme": "dark"})
        self.assertEqual(theme.status_code, 302)
        self.assertIn(b'data-theme="dark"', self.client.get("/").data)

        with self.app.test_client() as second_client:
            second_client.post(
                "/register",
                data={
                    "username": "second-user",
                    "password": "second-password-long",
                    "confirm_password": "second-password-long",
                },
            )
            self.assertIn(b'data-theme="light"', second_client.get("/").data)

    def test_workbook_export_has_multiple_sheets_and_omits_secrets(self):
        self.login()
        self.add_vehicle()
        with self.app.app_context():
            household = Household.query.one()
            household.discord_webhook_url = (
                "https://discord.com/api/webhooks/123456/super-secret-token"
            )
            owner = User.query.filter_by(username="owner").one()
            db.session.add(Vehicle(owner_id=owner.id, nickname="=2+2"))
            db.session.commit()

        workbook_response = self.client.get("/exports/workbook.xlsx")
        self.assertEqual(workbook_response.status_code, 200)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            workbook_response.content_type,
        )
        workbook = load_workbook(io.BytesIO(workbook_response.data), data_only=False)
        self.assertEqual(
            workbook.sheetnames,
            [
                "Garage",
                "Service History",
                "Mileage History",
                "Value History",
                "Reminders",
                "Household Members",
                "Settings",
            ],
        )
        garage = workbook["Garage"]
        self.assertIn("Bluebird", [cell.value for cell in garage["B"]])
        self.assertEqual(garage["B3"].value, "=2+2")
        self.assertEqual(garage["B3"].data_type, "s")
        self.assertIn("discord_notifications_configured", [
            cell.value for cell in workbook["Settings"][1]
        ])
        workbook_text = "\n".join(
            str(cell.value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        )
        self.assertNotIn("super-secret-token", workbook_text)

        backup = self.client.get("/exports/backup.zip")
        self.assertEqual(backup.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(backup.data)) as archive:
            self.assertIn("backup.json", archive.namelist())
            self.assertNotIn("password_hash", archive.read("backup.json").decode())

        with self.app.test_client() as anonymous_client:
            self.assertEqual(anonymous_client.get("/exports/workbook.xlsx").status_code, 302)
            self.assertEqual(anonymous_client.get("/exports/backup.zip").status_code, 302)

    def test_backup_restore_replaces_household_data_and_restores_receipts(self):
        self.login()
        with self.app.app_context():
            owner = User.query.filter_by(username="owner").one()
            household = owner.household
            household.name = "Restore Garage"
            household.discord_webhook_url = (
                "https://discord.com/api/webhooks/123456/restore-token"
            )
            owner.theme = "dark"
            vehicle = Vehicle(
                owner_id=owner.id,
                nickname="Restored Coupe",
                year=2022,
                make="Example",
                model="Roadster",
                vehicle_type="Coupe",
                vin="VIN-RESTORE",
                license_plate="RESTORE",
                odometer_entries=[
                    OdometerEntry(reading=32100, recorded_on=date(2026, 1, 2), note="Service")
                ],
                service_records=[
                    ServiceRecord(
                        title="Brake repair",
                        category="Repair",
                        completed_on=date(2026, 1, 3),
                        odometer=32100,
                        cost="425.50",
                        notes="Front brakes",
                        receipt_filename="brakes.pdf",
                        receipt_storage_name="original-receipt.pdf",
                    )
                ],
                value_entries=[
                    ValueEntry(
                        amount="12500.00",
                        valued_on=date(2026, 1, 4),
                        source="Estimate",
                        notes="Good condition",
                    )
                ],
                reminders=[
                    Reminder(
                        title="Rotate tires",
                        due_date=date(2026, 7, 1),
                        due_odometer=37000,
                        recurrence_days=180,
                        recurrence_miles=5000,
                        notes="",
                    )
                ],
            )
            db.session.add(vehicle)
            db.session.commit()
            receipt_path = Path(self.temp_dir.name) / "receipts" / "original-receipt.pdf"
            receipt_path.write_bytes(b"restorable receipt")

        backup = self.client.get("/exports/backup.zip")
        self.assertEqual(backup.status_code, 200)
        with self.app.app_context():
            owner = User.query.filter_by(username="owner").one()
            owner.theme = "light"
            owner.household.name = "Changed after backup"
            owner.household.discord_webhook_url = None
            db.session.add(Vehicle(owner_id=owner.id, nickname="Discard on restore"))
            db.session.commit()

        restored = self.client.post(
            "/exports/restore",
            data={
                "confirm_replace": "yes",
                "backup_file": (io.BytesIO(backup.data), "carfreeks-backup.zip"),
            },
            content_type="multipart/form-data",
            follow_redirects=True,
        )
        self.assertEqual(restored.status_code, 200)
        self.assertIn(b"1 vehicle(s) are now in your garage", restored.data)
        self.assertIn(b"Restored Coupe", restored.data)
        self.assertNotIn(b"Discard on restore", restored.data)

        with self.app.app_context():
            owner = User.query.filter_by(username="owner").one()
            self.assertEqual(owner.theme, "dark")
            self.assertEqual(owner.household.name, "Restore Garage")
            self.assertEqual(
                owner.household.discord_webhook_url,
                "https://discord.com/api/webhooks/123456/restore-token",
            )
            vehicles = Vehicle.query.all()
            self.assertEqual(len(vehicles), 1)
            self.assertEqual(vehicles[0].odometer_entries[0].reading, 32100)
            self.assertEqual(vehicles[0].service_records[0].cost, 425.50)
            self.assertEqual(vehicles[0].value_entries[0].source, "Estimate")
            self.assertEqual(vehicles[0].reminders[0].recurrence_miles, 5000)
            receipt_record = vehicles[0].service_records[0]
            self.assertNotEqual(receipt_record.receipt_storage_name, "original-receipt.pdf")
            restored_receipt = (
                Path(self.temp_dir.name)
                / "receipts"
                / receipt_record.receipt_storage_name
            )
            self.assertEqual(restored_receipt.read_bytes(), b"restorable receipt")

    def test_restore_requires_confirmation_and_rejects_invalid_backup(self):
        self.login()
        self.add_vehicle()
        with self.app.app_context():
            vehicle_id = Vehicle.query.one().id

        not_confirmed = self.client.post(
            "/exports/restore",
            data={"backup_file": (io.BytesIO(b"bad"), "bad.zip")},
            content_type="multipart/form-data",
            follow_redirects=True,
        )
        self.assertIn(b"Confirm garage replacement", not_confirmed.data)
        invalid = self.client.post(
            "/exports/restore",
            data={
                "confirm_replace": "yes",
                "backup_file": (io.BytesIO(b"not a zip"), "bad.zip"),
            },
            content_type="multipart/form-data",
            follow_redirects=True,
        )
        self.assertIn(b"valid CarFreeks backup ZIP", invalid.data)
        with self.app.app_context():
            self.assertEqual(Vehicle.query.one().id, vehicle_id)

    def test_discord_webhook_url_is_limited_to_discord_webhooks(self):
        self.assertTrue(
            valid_discord_webhook_url(
                "https://discord.com/api/webhooks/123456/test_token"
            )
        )
        self.assertFalse(valid_discord_webhook_url("https://example.com/api/webhooks/1/x"))
        self.assertFalse(valid_discord_webhook_url("http://discord.com/api/webhooks/1/x"))
        self.assertFalse(
            valid_discord_webhook_url(
                "https://discord.com.evil.example/api/webhooks/1/x"
            )
        )
        self.assertFalse(valid_discord_webhook_url("https://discord.com/api/webhooks/1/x?x=1"))

    def test_owner_can_save_and_test_discord_webhook(self):
        self.login()
        valid_url = "https://discord.com/api/webhooks/123456/test_token"
        saved = self.client.post(
            "/household",
            data={
                "name": "Garage Crew",
                "discord_webhook_url": valid_url,
            },
        )
        self.assertEqual(saved.status_code, 302)
        settings = self.client.get("/household")
        self.assertIn(b"Garage Crew", settings.data)
        self.assertNotIn(valid_url.encode(), settings.data)

        invalid = self.client.post(
            "/household",
            data={
                "name": "Garage Crew",
                "discord_webhook_url": "https://example.com/webhook",
            },
        )
        self.assertEqual(invalid.status_code, 200)
        self.assertIn(b"Enter a Discord webhook URL", invalid.data)

        with patch("notifications.send_discord_message") as send:
            response = self.client.post("/household/test-discord")
        self.assertEqual(response.status_code, 302)
        self.assertEqual(send.call_count, 1)
        self.assertIn("connection test", send.call_args.args[1])

    def test_discord_worker_notifies_each_due_reminder_once(self):
        from notifications import process_due_reminders

        self.login()
        self.add_vehicle()
        with self.app.app_context():
            vehicle = Vehicle.query.one()
            vehicle_id = vehicle.id
            household = Household.query.one()
            household.discord_webhook_url = (
                "https://discord.com/api/webhooks/123456/test_token"
            )
            reminder = Reminder(
                vehicle_id=vehicle_id,
                title="Oil change",
                due_date=date.today(),
            )
            db.session.add(reminder)
            db.session.commit()
            reminder_id = reminder.id

        with patch("notifications.send_discord_message") as send:
            self.assertEqual(process_due_reminders(self.app), 1)
            self.assertEqual(process_due_reminders(self.app), 0)
        self.assertEqual(send.call_count, 1)
        self.assertIn("Oil change", send.call_args.args[1])
        self.assertNotIn("@everyone", send.call_args.args[1])
        with self.app.app_context():
            self.assertIsNotNone(db.session.get(Reminder, reminder_id).last_notified_at)

    def test_discord_request_disables_mention_parsing(self):
        from notifications import send_discord_message

        with patch("notifications.build_opener") as build_opener:
            response_context = build_opener.return_value.open.return_value.__enter__.return_value
            response_context.status = 204
            send_discord_message(
                "https://discord.com/api/webhooks/123456/test_token",
                "CarFreeks reminder due",
            )
            request = build_opener.return_value.open.call_args.args[0]
        self.assertEqual(
            json.loads(request.data)["allowed_mentions"],
            {"parse": []},
        )

    def test_recurring_reminder_reschedules_on_completion(self):
        self.login()
        self.add_vehicle()
        with self.app.app_context():
            vehicle = Vehicle.query.one()
            vehicle_id = vehicle.id
            vehicle.odometer_entries.append(
                OdometerEntry(reading=14500, recorded_on=date.today())
            )
            reminder = Reminder(
                vehicle_id=vehicle_id,
                title="Oil change",
                due_date=date.today(),
                due_odometer=15000,
                recurrence_days=90,
                recurrence_miles=5000,
            )
            db.session.add(reminder)
            db.session.commit()
            reminder_id = reminder.id

        response = self.client.post(
            f"/vehicles/{vehicle_id}/reminders/{reminder_id}/toggle"
        )
        self.assertEqual(response.status_code, 302)
        self.assertIn(b"next recurring service is scheduled", self.client.get(f"/vehicles/{vehicle_id}").data)
        with self.app.app_context():
            reminder = db.session.get(Reminder, reminder_id)
            self.assertFalse(reminder.completed)
            self.assertEqual(reminder.due_date, date.today() + timedelta(days=90))
            self.assertEqual(reminder.due_odometer, 15000)
            self.assertIsNone(reminder.last_notified_at)

    def test_existing_owner_database_is_migrated_without_losing_login(self):
        legacy_database = Path(self.temp_dir.name) / "legacy.db"
        connection = sqlite3.connect(legacy_database)
        connection.execute(
            'CREATE TABLE "user" ('
            "id INTEGER PRIMARY KEY, "
            "username VARCHAR(64) UNIQUE NOT NULL, "
            "password_hash VARCHAR(256) NOT NULL)"
        )
        connection.execute(
            'INSERT INTO "user" (id, username, password_hash) VALUES (?, ?, ?)',
            (1, "legacy-owner", generate_password_hash("legacy-password-long")),
        )
        connection.execute(
            "CREATE TABLE reminder ("
            "id INTEGER PRIMARY KEY, vehicle_id INTEGER NOT NULL, "
            "title VARCHAR(120) NOT NULL, due_date DATE, due_odometer INTEGER, "
            "notes TEXT, completed BOOLEAN NOT NULL DEFAULT 0)"
        )
        connection.commit()
        connection.close()

        migration_app = create_app(
            {
                "TESTING": True,
                "SECRET_KEY": "migration-test-secret-key",
                "CARFREEKS_DATA_DIR": self.temp_dir.name,
                "UPLOAD_FOLDER": str(Path(self.temp_dir.name) / "receipts"),
                "SQLALCHEMY_DATABASE_URI": f"sqlite:///{legacy_database.as_posix()}",
                "WTF_CSRF_ENABLED": False,
            }
        )
        with migration_app.test_client() as legacy_client:
            login = legacy_client.post(
                "/login",
                data={
                    "username": "legacy-owner",
                    "password": "legacy-password-long",
                },
                follow_redirects=True,
            )
            self.assertIn(b"Your garage.", login.data)
            self.assertEqual(legacy_client.get("/household").status_code, 200)
        with migration_app.app_context():
            legacy_user = User.query.filter_by(username="legacy-owner").one()
            self.assertIsNotNone(legacy_user.household_id)
            db.session.remove()
            db.engine.dispose()

    def add_vehicle(self):
        return self.client.post(
            "/vehicles/new",
            data={
                "nickname": "Bluebird",
                "year": "2020",
                "make": "Example",
                "model": "Commuter",
                "vehicle_type": "Car",
                "vin": "",
                "license_plate": "",
            },
        )

    def test_owner_login_and_vehicle_dashboard(self):
        self.assertEqual(self.client.get("/").status_code, 302)
        self.login()

        response = self.add_vehicle()

        self.assertEqual(response.status_code, 302)
        dashboard = self.client.get("/")
        self.assertEqual(dashboard.status_code, 200)
        self.assertIn(b"Bluebird", dashboard.data)
        self.assertIn(b"2020 Example Commuter", dashboard.data)

    def test_state_changes_require_csrf_tokens(self):
        self.login()
        self.app.config["WTF_CSRF_ENABLED"] = True

        response = self.add_vehicle()

        self.assertEqual(response.status_code, 400)
        with self.app.app_context():
            self.assertEqual(Vehicle.query.count(), 0)

    def test_vehicle_records_and_receipt_download(self):
        self.login()
        self.add_vehicle()
        with self.app.app_context():
            vehicle = Vehicle.query.one()
            vehicle_id = vehicle.id

        self.client.post(
            f"/vehicles/{vehicle_id}/odometer/new",
            data={"reading": "32500", "recorded_on": date.today().isoformat(), "note": ""},
        )
        self.client.post(
            f"/vehicles/{vehicle_id}/value/new",
            data={
                "amount": "18750.00",
                "valued_on": date.today().isoformat(),
                "source": "Owner estimate",
                "notes": "",
            },
        )
        response = self.client.post(
            f"/vehicles/{vehicle_id}/service/new",
            data={
                "title": "Oil change",
                "category": "Maintenance",
                "completed_on": date.today().isoformat(),
                "odometer": "32500",
                "cost": "64.95",
                "notes": "Synthetic oil",
                "receipt": (io.BytesIO(b"test receipt"), "receipt.pdf"),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 302)

        detail = self.client.get(f"/vehicles/{vehicle_id}")
        self.assertEqual(detail.status_code, 200)
        self.assertIn(b"32,500", detail.data)
        self.assertIn(b"$18,750.00", detail.data)
        self.assertIn(b"Oil change", detail.data)

        with self.app.app_context():
            record_id = Vehicle.query.one().service_records[0].id
        receipt = self.client.get(f"/receipts/{record_id}")
        try:
            self.assertEqual(receipt.status_code, 200)
            self.assertIn("attachment", receipt.headers["Content-Disposition"])
            self.assertEqual(receipt.data, b"test receipt")
        finally:
            receipt.close()

    def test_receipts_and_vehicle_records_are_owner_scoped(self):
        self.login()
        self.add_vehicle()
        with self.app.app_context():
            owner = User(username="different-owner")
            owner.set_password("different-password-long")
            db.session.add(owner)
            db.session.flush()
            vehicle = Vehicle(owner_id=owner.id, nickname="Private car")
            db.session.add(vehicle)
            db.session.commit()
            vehicle_id = vehicle.id
            record = ServiceRecord(
                vehicle_id=vehicle_id,
                title="Private repair",
                category="Repair",
                completed_on=date.today(),
                cost=0,
                receipt_filename="private.pdf",
                receipt_storage_name="private.pdf",
            )
            db.session.add(record)
            db.session.commit()
            record_id = record.id

        self.assertEqual(self.client.get(f"/vehicles/{vehicle_id}").status_code, 404)
        self.assertEqual(self.client.get(f"/receipts/{record_id}").status_code, 404)

    def test_vehicle_can_be_updated_and_deleted(self):
        self.login()
        self.add_vehicle()
        with self.app.app_context():
            vehicle_id = Vehicle.query.one().id

        update = self.client.post(
            f"/vehicles/{vehicle_id}/edit",
            data={
                "nickname": "Bluebird GT",
                "year": "2020",
                "make": "Example",
                "model": "Commuter",
                "vehicle_type": "Car",
                "vin": "",
                "license_plate": "ABC123",
            },
        )
        self.assertEqual(update.status_code, 302)
        self.assertIn(b"Bluebird GT", self.client.get("/").data)

        delete = self.client.post(f"/vehicles/{vehicle_id}/delete")
        self.assertEqual(delete.status_code, 302)
        dashboard = self.client.get("/")
        self.assertIn(b"0 in your garage", dashboard.data)
        self.assertIn(b"Start your garage", dashboard.data)

    def test_service_rejects_unsupported_receipt_types(self):
        self.login()
        self.add_vehicle()
        with self.app.app_context():
            vehicle_id = Vehicle.query.one().id

        response = self.client.post(
            f"/vehicles/{vehicle_id}/service/new",
            data={
                "title": "Oil change",
                "category": "Maintenance",
                "completed_on": date.today().isoformat(),
                "odometer": "",
                "cost": "40",
                "notes": "",
                "receipt": (io.BytesIO(b"not an image"), "receipt.exe"),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Use a PDF or supported image", response.data)

    def test_reminders_require_a_date_or_mileage(self):
        self.login()
        self.add_vehicle()
        with self.app.app_context():
            vehicle_id = Vehicle.query.one().id

        invalid = self.client.post(
            f"/vehicles/{vehicle_id}/reminders/new",
            data={"title": "Tire rotation", "due_date": "", "due_odometer": ""},
        )
        self.assertEqual(invalid.status_code, 200)
        self.assertIn(b"Enter a first due date", invalid.data)

        valid = self.client.post(
            f"/vehicles/{vehicle_id}/reminders/new",
            data={
                "title": "Tire rotation",
                "due_date": (date.today() + timedelta(days=10)).isoformat(),
                "due_odometer": "35000",
                "recurrence_miles": "5000",
                "notes": "",
            },
        )
        self.assertEqual(valid.status_code, 302)

    def test_reminder_status_and_latest_manual_value(self):
        today = date.today()
        overdue = Reminder(title="Oil", due_date=today - timedelta(days=1))
        due_soon = Reminder(title="Tires", due_date=today + timedelta(days=10))
        mileage_due = Reminder(title="Brakes", due_odometer=33500)
        upcoming = Reminder(title="Inspection", due_date=today + timedelta(days=60))

        self.assertEqual(reminder_status(overdue, 30000), "Overdue")
        self.assertEqual(reminder_status(due_soon, 30000), "Due soon")
        self.assertEqual(reminder_status(mileage_due, 33000), "Due soon")
        self.assertEqual(reminder_status(upcoming, 30000), "Upcoming")

        earlier = ValueEntry(amount=10000, valued_on=today - timedelta(days=30))
        latest = ValueEntry(amount=9000, valued_on=today)
        vehicle = Vehicle(value_entries=[earlier, latest])
        self.assertIs(vehicle.latest_value, latest)


if __name__ == "__main__":
    unittest.main()
