import hashlib
import io
import json
import logging
import os
import re
import secrets
import uuid
import zipfile
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.styles import Font
from flask import (
    Flask,
    abort,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    session,
    send_file,
    send_from_directory,
    url_for,
)
from flask_login import (
    LoginManager,
    UserMixin,
    current_user,
    login_required,
    login_user,
    logout_user,
)
from flask_sqlalchemy import SQLAlchemy
from flask_wtf import FlaskForm, CSRFProtect
from sqlalchemy import Numeric, inspect
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename
from wtforms import (
    DateField,
    DecimalField,
    BooleanField,
    FileField,
    HiddenField,
    IntegerField,
    PasswordField,
    SelectField,
    StringField,
    SubmitField,
    TextAreaField,
)
from wtforms.validators import (
    DataRequired,
    EqualTo,
    InputRequired,
    Length,
    NumberRange,
    Optional,
)


BASE_DIR = Path(__file__).resolve().parent
ALLOWED_RECEIPT_EXTENSIONS = {".pdf", ".jpg", ".jpeg", ".png", ".webp", ".heic"}
DISCORD_WEBHOOK_PATTERN = re.compile(r"^/api/webhooks/[0-9]+/[A-Za-z0-9._-]+$")
MAX_RECEIPT_SIZE = 10 * 1024 * 1024
MAX_BACKUP_SIZE = 100 * 1024 * 1024
MAX_BACKUP_JSON_SIZE = 10 * 1024 * 1024
logger = logging.getLogger(__name__)


def utc_now():
    return datetime.now(timezone.utc).replace(tzinfo=None)

db = SQLAlchemy()
login_manager = LoginManager()
csrf = CSRFProtect()


class User(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(64), unique=True, nullable=False)
    password_hash = db.Column(db.String(256), nullable=False)
    household_id = db.Column(db.Integer, db.ForeignKey("household.id"), nullable=True)
    role = db.Column(db.String(16), nullable=False, default="owner")
    theme = db.Column(db.String(8), nullable=False, default="light")
    household = db.relationship("Household", back_populates="members")

    def set_password(self, password):
        self.password_hash = generate_password_hash(password)

    def check_password(self, password):
        return check_password_hash(self.password_hash, password)


class Household(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), nullable=False)
    discord_webhook_url = db.Column(db.String(512))
    members = db.relationship("User", back_populates="household")


class HouseholdInvitation(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    household_id = db.Column(db.Integer, db.ForeignKey("household.id"), nullable=False)
    created_by_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    token_hash = db.Column(db.String(64), nullable=False, unique=True)
    created_at = db.Column(db.DateTime, nullable=False)
    expires_at = db.Column(db.DateTime, nullable=False)
    accepted_at = db.Column(db.DateTime)
    household = db.relationship("Household")


class Vehicle(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    owner_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    nickname = db.Column(db.String(80), nullable=False)
    year = db.Column(db.Integer)
    make = db.Column(db.String(80))
    model = db.Column(db.String(80))
    vehicle_type = db.Column(db.String(32), nullable=False, default="Car")
    vin = db.Column(db.String(32))
    license_plate = db.Column(db.String(24))
    created_at = db.Column(
        db.DateTime(timezone=True),
        nullable=False,
        default=lambda: datetime.now(timezone.utc),
    )
    odometer_entries = db.relationship(
        "OdometerEntry", backref="vehicle", cascade="all, delete-orphan"
    )
    service_records = db.relationship(
        "ServiceRecord", backref="vehicle", cascade="all, delete-orphan"
    )
    value_entries = db.relationship(
        "ValueEntry", backref="vehicle", cascade="all, delete-orphan"
    )
    reminders = db.relationship(
        "Reminder", backref="vehicle", cascade="all, delete-orphan"
    )

    @property
    def current_odometer(self):
        latest = max(
            self.odometer_entries,
            key=lambda entry: (entry.recorded_on, entry.id or 0),
            default=None,
        )
        return latest.reading if latest else None

    @property
    def latest_value(self):
        return max(
            self.value_entries,
            key=lambda entry: (entry.valued_on, entry.id or 0),
            default=None,
        )


class OdometerEntry(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    vehicle_id = db.Column(db.Integer, db.ForeignKey("vehicle.id"), nullable=False)
    reading = db.Column(db.Integer, nullable=False)
    recorded_on = db.Column(db.Date, nullable=False, default=date.today)
    note = db.Column(db.String(250))


class ServiceRecord(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    vehicle_id = db.Column(db.Integer, db.ForeignKey("vehicle.id"), nullable=False)
    title = db.Column(db.String(120), nullable=False)
    category = db.Column(db.String(24), nullable=False)
    completed_on = db.Column(db.Date, nullable=False, default=date.today)
    odometer = db.Column(db.Integer)
    cost = db.Column(Numeric(10, 2), nullable=False, default=Decimal("0.00"))
    notes = db.Column(db.Text)
    receipt_filename = db.Column(db.String(255))
    receipt_storage_name = db.Column(db.String(64))


class ValueEntry(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    vehicle_id = db.Column(db.Integer, db.ForeignKey("vehicle.id"), nullable=False)
    amount = db.Column(Numeric(10, 2), nullable=False)
    valued_on = db.Column(db.Date, nullable=False, default=date.today)
    source = db.Column(db.String(120))
    notes = db.Column(db.Text)


class Reminder(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    vehicle_id = db.Column(db.Integer, db.ForeignKey("vehicle.id"), nullable=False)
    title = db.Column(db.String(120), nullable=False)
    due_date = db.Column(db.Date)
    due_odometer = db.Column(db.Integer)
    recurrence_days = db.Column(db.Integer)
    recurrence_miles = db.Column(db.Integer)
    notes = db.Column(db.Text)
    completed = db.Column(db.Boolean, nullable=False, default=False)
    last_notified_at = db.Column(db.DateTime)


class LoginForm(FlaskForm):
    username = StringField("Username", validators=[DataRequired(), Length(max=64)])
    password = PasswordField("Password", validators=[DataRequired()])
    submit = SubmitField("Sign in")


class AccountSetupForm(FlaskForm):
    invite_token = HiddenField()
    username = StringField("Username", validators=[DataRequired(), Length(max=64)])
    password = PasswordField(
        "Password", validators=[DataRequired(), Length(min=12, max=128)]
    )
    confirm_password = PasswordField(
        "Confirm password",
        validators=[DataRequired(), EqualTo("password", message="Passwords must match.")],
    )
    submit = SubmitField("Create my account")


class AccountForm(FlaskForm):
    current_password = PasswordField("Current password", validators=[DataRequired()])
    username = StringField("Username", validators=[DataRequired(), Length(max=64)])
    password = PasswordField(
        "New password",
        validators=[Optional(), Length(min=12, max=128)],
        description="Leave blank to keep your current password.",
    )
    confirm_password = PasswordField(
        "Confirm new password",
        validators=[
            EqualTo("password", message="New passwords must match."),
        ],
    )
    submit = SubmitField("Save account")


class HouseholdForm(FlaskForm):
    name = StringField("Household name", validators=[DataRequired(), Length(max=100)])
    discord_webhook_url = StringField(
        "Discord webhook URL", validators=[Optional(), Length(max=512)]
    )
    remove_discord_webhook = BooleanField("Remove saved Discord webhook")
    submit = SubmitField("Save household")


class ReminderForm(FlaskForm):
    title = StringField("Maintenance item", validators=[DataRequired(), Length(max=120)])
    due_date = DateField("First due date", validators=[Optional()])
    due_odometer = IntegerField(
        "First due at odometer reading", validators=[Optional(), NumberRange(min=0)]
    )
    recurrence_days = IntegerField(
        "Repeat every (days)", validators=[Optional(), NumberRange(min=1, max=36500)]
    )
    recurrence_miles = IntegerField(
        "Repeat every (miles)", validators=[Optional(), NumberRange(min=1)]
    )
    notes = TextAreaField("Notes", validators=[Optional(), Length(max=4000)])
    submit = SubmitField("Save reminder")

    def validate(self, extra_validators=None):
        valid = super().validate(extra_validators=extra_validators)
        if self.due_date.data is None and self.due_odometer.data is None:
            self.due_date.errors.append("Enter a first due date, an odometer reading, or both.")
            return False
        if self.recurrence_days.data is not None and self.due_date.data is None:
            self.recurrence_days.errors.append("Set a first due date for a date-based repeat.")
            return False
        if self.recurrence_miles.data is not None and self.due_odometer.data is None:
            self.recurrence_miles.errors.append(
                "Set a first due odometer reading for a mileage-based repeat."
            )
            return False
        return valid


class VehicleForm(FlaskForm):
    nickname = StringField("Vehicle name", validators=[DataRequired(), Length(max=80)])
    year = IntegerField("Year", validators=[Optional(), NumberRange(min=1886, max=2100)])
    make = StringField("Make", validators=[Optional(), Length(max=80)])
    model = StringField("Model", validators=[Optional(), Length(max=80)])
    vehicle_type = SelectField(
        "Type",
        choices=[
            ("Car", "Car"),
            ("SUV", "SUV"),
            ("Truck", "Truck"),
            ("Van", "Van"),
            ("Motorcycle", "Motorcycle"),
            ("Other", "Other"),
        ],
        validators=[DataRequired()],
    )
    vin = StringField("VIN", validators=[Optional(), Length(max=32)])
    license_plate = StringField("License plate", validators=[Optional(), Length(max=24)])
    submit = SubmitField("Save vehicle")


class OdometerForm(FlaskForm):
    reading = IntegerField(
        "Odometer reading", validators=[InputRequired(), NumberRange(min=0)]
    )
    recorded_on = DateField("Reading date", validators=[InputRequired()], default=date.today)
    note = StringField("Note", validators=[Optional(), Length(max=250)])
    submit = SubmitField("Add reading")


class ServiceForm(FlaskForm):
    title = StringField("Service or repair", validators=[DataRequired(), Length(max=120)])
    category = SelectField(
        "Category",
        choices=[("Maintenance", "Maintenance"), ("Repair", "Repair")],
        validators=[DataRequired()],
    )
    completed_on = DateField("Date completed", validators=[InputRequired()], default=date.today)
    odometer = IntegerField(
        "Odometer reading", validators=[Optional(), NumberRange(min=0)]
    )
    cost = DecimalField(
        "Cost", places=2, validators=[InputRequired(), NumberRange(min=Decimal("0"))]
    )
    notes = TextAreaField("Notes", validators=[Optional(), Length(max=4000)])
    receipt = FileField("Receipt (PDF or image)")
    submit = SubmitField("Save record")


class ValueForm(FlaskForm):
    amount = DecimalField(
        "Estimated value",
        places=2,
        validators=[InputRequired(), NumberRange(min=Decimal("0.01"))],
    )
    valued_on = DateField("Value date", validators=[InputRequired()], default=date.today)
    source = StringField("Source", validators=[Optional(), Length(max=120)])
    notes = TextAreaField("Notes", validators=[Optional(), Length(max=4000)])
    submit = SubmitField("Save value")


@login_manager.user_loader
def load_user(user_id):
    return db.session.get(User, int(user_id))


def migrate_existing_records():
    inspector = inspect(db.engine)
    if "user" in inspector.get_table_names():
        user_columns = {column["name"] for column in inspector.get_columns("user")}
        quote = db.engine.dialect.identifier_preparer.quote
        with db.engine.begin() as connection:
            if "household_id" not in user_columns:
                connection.exec_driver_sql(
                    f"ALTER TABLE {quote('user')} ADD COLUMN household_id "
                    "INTEGER REFERENCES household (id)"
                )
            if "role" not in user_columns:
                connection.exec_driver_sql(
                    f"ALTER TABLE {quote('user')} ADD COLUMN role "
                    "VARCHAR(16) NOT NULL DEFAULT 'owner'"
                )
            if "theme" not in user_columns:
                connection.exec_driver_sql(
                    f"ALTER TABLE {quote('user')} ADD COLUMN theme "
                    "VARCHAR(8) NOT NULL DEFAULT 'light'"
                )

    if "reminder" in inspect(db.engine).get_table_names():
        reminder_columns = {
            column["name"] for column in inspect(db.engine).get_columns("reminder")
        }
        quote = db.engine.dialect.identifier_preparer.quote
        with db.engine.begin() as connection:
            for name, column_type in (
                ("recurrence_days", "INTEGER"),
                ("recurrence_miles", "INTEGER"),
                ("last_notified_at", "DATETIME"),
            ):
                if name not in reminder_columns:
                    connection.exec_driver_sql(
                        f"ALTER TABLE {quote('reminder')} ADD COLUMN "
                        f"{quote(name)} {column_type}"
                    )

    for user in User.query.filter(User.household_id.is_(None)).all():
        household = Household(name=f"{user.username}'s Garage")
        db.session.add(household)
        db.session.flush()
        user.household_id = household.id
        user.role = "owner"
    db.session.commit()


def invitation_token_hash(token):
    return hashlib.sha256(token.encode("utf-8")).hexdigest()


def valid_discord_webhook_url(value):
    if not value:
        return True
    try:
        parsed = urlsplit(value.strip())
        return (
            parsed.scheme == "https"
            and parsed.hostname in {"discord.com", "discordapp.com"}
            and parsed.port in {None, 443}
            and parsed.username is None
            and parsed.password is None
            and parsed.query == ""
            and parsed.fragment == ""
            and DISCORD_WEBHOOK_PATTERN.fullmatch(parsed.path) is not None
        )
    except ValueError:
        return False


def is_household_owner(user):
    return user.role == "owner" and user.household_id is not None


def create_app(test_config=None):
    app = Flask(__name__)
    data_dir = Path(os.environ.get("CARFREEKS_DATA_DIR", BASE_DIR / "data"))
    database_url = os.environ.get("CARFREEKS_DATABASE_URL")
    app.config.from_mapping(
        SECRET_KEY=os.environ.get("CARFREEKS_SECRET_KEY"),
        SQLALCHEMY_DATABASE_URI=database_url or f"sqlite:///{data_dir / 'carfreeks.db'}",
        SQLALCHEMY_TRACK_MODIFICATIONS=False,
        SQLALCHEMY_ENGINE_OPTIONS={"connect_args": {"check_same_thread": False}},
        CARFREEKS_DATA_DIR=str(data_dir),
        UPLOAD_FOLDER=str(data_dir / "receipts"),
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
        SESSION_COOKIE_SECURE=os.environ.get("CARFREEKS_COOKIE_SECURE", "").lower()
        in {"1", "true", "yes"},
        MAX_CONTENT_LENGTH=MAX_BACKUP_SIZE,
    )
    if test_config:
        app.config.update(test_config)

    if not app.config["SECRET_KEY"]:
        raise RuntimeError("Set CARFREEKS_SECRET_KEY to a long, random secret.")

    Path(app.config["CARFREEKS_DATA_DIR"]).mkdir(parents=True, exist_ok=True)
    Path(app.config["UPLOAD_FOLDER"]).mkdir(parents=True, exist_ok=True)

    db.init_app(app)
    login_manager.init_app(app)
    login_manager.login_view = "login"
    login_manager.login_message = "Sign in to view your vehicles."
    login_manager.login_message_category = "info"
    csrf.init_app(app)

    with app.app_context():
        db.create_all()
        migrate_existing_records()

    register_routes(app)
    return app


def get_vehicle_or_404(vehicle_id):
    member_ids = db.select(User.id).where(User.household_id == current_user.household_id)
    return Vehicle.query.filter(
        Vehicle.id == vehicle_id, Vehicle.owner_id.in_(member_ids)
    ).first_or_404()


def export_datasets():
    member_ids = db.select(User.id).where(
        User.household_id == current_user.household_id
    )
    members = User.query.filter_by(household_id=current_user.household_id).order_by(
        User.id
    ).all()
    vehicles = Vehicle.query.filter(Vehicle.owner_id.in_(member_ids)).order_by(
        Vehicle.id
    ).all()
    owners = {member.id: member.username for member in members}
    household = current_user.household
    exports = {
        "garage": [[
            "vehicle_id", "vehicle_name", "type", "year", "make", "model", "vin",
            "license_plate", "added_at", "added_by", "current_odometer", "latest_estimated_value",
        ]],
        "service-history": [[
            "vehicle_id", "vehicle_name", "service_record_id", "title", "category",
            "completed_on", "odometer", "cost", "notes", "receipt_filename",
        ]],
        "mileage-history": [[
            "vehicle_id", "vehicle_name", "recorded_on", "odometer", "note",
        ]],
        "value-history": [[
            "vehicle_id", "vehicle_name", "valued_on", "amount", "source", "notes",
        ]],
        "reminders": [[
            "vehicle_id", "vehicle_name", "title", "due_date", "due_odometer",
            "recurrence_days", "recurrence_miles", "completed", "notes",
        ]],
        "household-members": [["user_id", "username", "role", "theme"]],
        "settings": [[
            "household_name", "current_user", "theme", "discord_notifications_configured",
        ], [
            household.name if household else "",
            current_user.username,
            current_user.theme,
            bool(household and household.discord_webhook_url),
        ]],
    }
    for member in members:
        exports["household-members"].append(
            [member.id, member.username, member.role, member.theme]
        )
    for vehicle in vehicles:
        exports["garage"].append([
            vehicle.id,
            vehicle.nickname,
            vehicle.vehicle_type,
            vehicle.year,
            vehicle.make,
            vehicle.model,
            vehicle.vin,
            vehicle.license_plate,
            vehicle.created_at.isoformat() if vehicle.created_at else "",
            owners.get(vehicle.owner_id, ""),
            vehicle.current_odometer if vehicle.current_odometer is not None else "",
            vehicle.latest_value.amount if vehicle.latest_value else "",
        ])
        for record in vehicle.service_records:
            exports["service-history"].append([
                vehicle.id, vehicle.nickname, record.id, record.title, record.category,
                record.completed_on.isoformat() if record.completed_on else "",
                record.odometer if record.odometer is not None else "",
                record.cost, record.notes or "", record.receipt_filename or "",
            ])
        for entry in vehicle.odometer_entries:
            exports["mileage-history"].append([
                vehicle.id, vehicle.nickname, entry.recorded_on.isoformat(),
                entry.reading, entry.note or "",
            ])
        for entry in vehicle.value_entries:
            exports["value-history"].append([
                vehicle.id, vehicle.nickname, entry.valued_on.isoformat(),
                entry.amount, entry.source or "", entry.notes or "",
            ])
        for reminder in vehicle.reminders:
            exports["reminders"].append([
                vehicle.id, vehicle.nickname, reminder.title,
                reminder.due_date.isoformat() if reminder.due_date else "",
                reminder.due_odometer if reminder.due_odometer is not None else "",
                reminder.recurrence_days if reminder.recurrence_days is not None else "",
                reminder.recurrence_miles if reminder.recurrence_miles is not None else "",
                reminder.completed, reminder.notes or "",
            ])
    return exports


def create_workbook():
    workbook = Workbook()
    workbook.remove(workbook.active)
    titles = {
        "garage": "Garage",
        "service-history": "Service History",
        "mileage-history": "Mileage History",
        "value-history": "Value History",
        "reminders": "Reminders",
        "household-members": "Household Members",
        "settings": "Settings",
    }
    for dataset, rows in export_datasets().items():
        sheet = workbook.create_sheet(titles[dataset])
        for row in rows:
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 42)
            sheet.column_dimensions[column[0].column_letter].width = max(width, 12)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


class BackupValidationError(ValueError):
    pass


def _backup_text(value, name, max_length, *, required=True):
    if value is None and not required:
        return None
    if not isinstance(value, str) or len(value) > max_length:
        raise BackupValidationError(f"Invalid {name} in backup.")
    if required and not value.strip():
        raise BackupValidationError(f"Missing {name} in backup.")
    return value


def _backup_integer(value, name, *, required=False, minimum=0, maximum=2**31 - 1):
    if value is None and not required:
        return None
    if isinstance(value, bool) or not isinstance(value, int):
        raise BackupValidationError(f"Invalid {name} in backup.")
    if value < minimum or value > maximum:
        raise BackupValidationError(f"Invalid {name} in backup.")
    return value


def _backup_date(value, name, *, required=False):
    if value is None and not required:
        return None
    if not isinstance(value, str):
        raise BackupValidationError(f"Invalid {name} in backup.")
    try:
        return date.fromisoformat(value)
    except ValueError as error:
        raise BackupValidationError(f"Invalid {name} in backup.") from error


def _backup_decimal(value, name):
    if isinstance(value, bool) or not isinstance(value, (str, int, float)):
        raise BackupValidationError(f"Invalid {name} in backup.")
    try:
        result = Decimal(str(value))
    except Exception as error:
        raise BackupValidationError(f"Invalid {name} in backup.") from error
    if not result.is_finite() or result < 0 or result > Decimal("99999999.99"):
        raise BackupValidationError(f"Invalid {name} in backup.")
    return result


def create_backup_archive(upload_folder):
    member_ids = db.select(User.id).where(
        User.household_id == current_user.household_id
    )
    vehicles = Vehicle.query.filter(Vehicle.owner_id.in_(member_ids)).order_by(
        Vehicle.id
    ).all()
    receipts = {}
    backup = {
        "format": "carfreeks-backup",
        "version": 1,
        "household": {
            "name": current_user.household.name,
            "discord_webhook_url": current_user.household.discord_webhook_url,
        },
        "theme": current_user.theme,
        "vehicles": [],
    }

    for vehicle in vehicles:
        saved_vehicle = {
            "nickname": vehicle.nickname,
            "year": vehicle.year,
            "make": vehicle.make,
            "model": vehicle.model,
            "vehicle_type": vehicle.vehicle_type,
            "vin": vehicle.vin,
            "license_plate": vehicle.license_plate,
            "created_at": vehicle.created_at.isoformat() if vehicle.created_at else None,
            "odometer_entries": [],
            "service_records": [],
            "value_entries": [],
            "reminders": [],
        }
        for entry in vehicle.odometer_entries:
            saved_vehicle["odometer_entries"].append({
                "reading": entry.reading,
                "recorded_on": entry.recorded_on.isoformat(),
                "note": entry.note,
            })
        for record in vehicle.service_records:
            saved_record = {
                "title": record.title,
                "category": record.category,
                "completed_on": record.completed_on.isoformat(),
                "odometer": record.odometer,
                "cost": str(record.cost),
                "notes": record.notes,
                "receipt_filename": record.receipt_filename,
                "receipt_path": None,
            }
            if record.receipt_storage_name:
                stored_name = record.receipt_storage_name
                receipt_path = Path(upload_folder) / stored_name
                resolved_upload_folder = Path(upload_folder).resolve()
                if (
                    Path(stored_name).name != stored_name
                    or receipt_path.suffix.lower() not in ALLOWED_RECEIPT_EXTENSIONS
                    or not receipt_path.resolve().is_relative_to(resolved_upload_folder)
                    or not receipt_path.is_file()
                    or receipt_path.stat().st_size > MAX_RECEIPT_SIZE
                ):
                    raise OSError("A saved receipt is missing or invalid.")
                archive_name = f"receipts/receipt-{len(receipts) + 1:06d}{receipt_path.suffix.lower()}"
                receipts[archive_name] = receipt_path.read_bytes()
                saved_record["receipt_path"] = archive_name
            saved_vehicle["service_records"].append(saved_record)
        for entry in vehicle.value_entries:
            saved_vehicle["value_entries"].append({
                "amount": str(entry.amount),
                "valued_on": entry.valued_on.isoformat(),
                "source": entry.source,
                "notes": entry.notes,
            })
        for reminder in vehicle.reminders:
            saved_vehicle["reminders"].append({
                "title": reminder.title,
                "due_date": reminder.due_date.isoformat() if reminder.due_date else None,
                "due_odometer": reminder.due_odometer,
                "recurrence_days": reminder.recurrence_days,
                "recurrence_miles": reminder.recurrence_miles,
                "notes": reminder.notes,
                "completed": reminder.completed,
                "last_notified_at": (
                    reminder.last_notified_at.isoformat()
                    if reminder.last_notified_at
                    else None
                ),
            })
        backup["vehicles"].append(saved_vehicle)

    payload = io.BytesIO()
    with zipfile.ZipFile(payload, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "manifest.json",
            json.dumps({
                "format": "carfreeks-backup",
                "version": 1,
                "created_at": utc_now().isoformat(),
            }),
        )
        archive.writestr("backup.json", json.dumps(backup, ensure_ascii=False))
        for path, content in receipts.items():
            archive.writestr(path, content)
    payload.seek(0)
    return payload


def read_backup_archive(upload):
    try:
        archive = zipfile.ZipFile(upload)
    except zipfile.BadZipFile as error:
        raise BackupValidationError("Choose a valid CarFreeks backup ZIP file.") from error
    with archive:
        members = archive.infolist()
        if (
            len(members) > 50002
            or sum(member.file_size for member in members) > MAX_BACKUP_SIZE
            or any(member.flag_bits & 0x1 for member in members)
            or any(
                member.is_dir()
                or member.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}
                for member in members
            )
        ):
            raise BackupValidationError("The backup exceeds supported size limits.")
        names = [member.filename for member in members]
        if len(names) != len(set(names)):
            raise BackupValidationError("The backup contains duplicate files.")
        if any(
            name not in {"manifest.json", "backup.json"}
            and re.fullmatch(
                r"receipts/receipt-[0-9]{6}\.(?:pdf|jpg|jpeg|png|webp|heic)",
                name,
            ) is None
            for name in names
        ):
            raise BackupValidationError("The backup contains an unsupported file.")
        if not {"manifest.json", "backup.json"}.issubset(names):
            raise BackupValidationError("This is not a complete CarFreeks backup.")
        if archive.testzip() is not None:
            raise BackupValidationError("The backup archive is damaged.")
        manifest_content = archive.read("manifest.json")
        backup_content = archive.read("backup.json")
        if len(manifest_content) > 4096 or len(backup_content) > MAX_BACKUP_JSON_SIZE:
            raise BackupValidationError("The backup contains oversized data.")
        try:
            manifest = json.loads(manifest_content)
            backup = json.loads(backup_content)
        except (UnicodeDecodeError, json.JSONDecodeError) as error:
            raise BackupValidationError("The backup data cannot be read.") from error

        if (
            not isinstance(manifest, dict)
            or manifest.get("format") != "carfreeks-backup"
            or manifest.get("version") != 1
            or not isinstance(backup, dict)
            or backup.get("format") != "carfreeks-backup"
            or backup.get("version") != 1
        ):
            raise BackupValidationError("This CarFreeks backup version is unsupported.")
        household = backup.get("household")
        vehicles = backup.get("vehicles")
        if not isinstance(household, dict) or not isinstance(vehicles, list) or len(vehicles) > 10000:
            raise BackupValidationError("The backup data has an invalid structure.")
        name = _backup_text(household.get("name"), "household name", 100)
        webhook_url = household.get("discord_webhook_url")
        if webhook_url is not None and (
            not isinstance(webhook_url, str)
            or len(webhook_url) > 512
            or not valid_discord_webhook_url(webhook_url)
        ):
            raise BackupValidationError("The backup contains an invalid Discord webhook.")
        theme = backup.get("theme")
        if not isinstance(theme, str) or theme not in {"light", "dark"}:
            raise BackupValidationError("The backup contains an invalid theme.")

        expected_receipts = set()
        parsed_vehicles = []
        total_records = 0
        for vehicle_data in vehicles:
            if not isinstance(vehicle_data, dict):
                raise BackupValidationError("The backup contains an invalid vehicle.")
            nickname = _backup_text(vehicle_data.get("nickname"), "vehicle name", 80)
            vehicle_type = _backup_text(
                vehicle_data.get("vehicle_type"), "vehicle type", 32
            )
            year = _backup_integer(vehicle_data.get("year"), "vehicle year", maximum=9999)
            if year is not None and year < 1886:
                raise BackupValidationError("Invalid vehicle year in backup.")
            fields = {
                "nickname": nickname,
                "year": year,
                "make": _backup_text(vehicle_data.get("make"), "vehicle make", 80, required=False),
                "model": _backup_text(vehicle_data.get("model"), "vehicle model", 80, required=False),
                "vehicle_type": vehicle_type,
                "vin": _backup_text(vehicle_data.get("vin"), "VIN", 32, required=False),
                "license_plate": _backup_text(
                    vehicle_data.get("license_plate"), "license plate", 24, required=False
                ),
            }
            created_at = vehicle_data.get("created_at")
            if created_at is not None:
                if not isinstance(created_at, str):
                    raise BackupValidationError("Invalid vehicle creation date in backup.")
                try:
                    fields["created_at"] = datetime.fromisoformat(created_at)
                except ValueError as error:
                    raise BackupValidationError(
                        "Invalid vehicle creation date in backup."
                    ) from error

            records = {}
            for collection in (
                "odometer_entries", "service_records", "value_entries", "reminders"
            ):
                values = vehicle_data.get(collection)
                if not isinstance(values, list):
                    raise BackupValidationError("The backup contains invalid vehicle records.")
                total_records += len(values)
                if total_records > 50000:
                    raise BackupValidationError("The backup contains too many records.")
                records[collection] = values
            parsed_vehicles.append((fields, records))

        found_receipts = set(names) - {"manifest.json", "backup.json"}
        referenced_receipts = set()
        validated = []
        for fields, records in parsed_vehicles:
            parsed = {"fields": fields}
            parsed["odometer_entries"] = []
            for entry in records["odometer_entries"]:
                if not isinstance(entry, dict):
                    raise BackupValidationError("Invalid mileage record in backup.")
                parsed["odometer_entries"].append({
                    "reading": _backup_integer(entry.get("reading"), "mileage", required=True),
                    "recorded_on": _backup_date(entry.get("recorded_on"), "mileage date", required=True),
                    "note": _backup_text(entry.get("note"), "mileage note", 250, required=False),
                })
            parsed["service_records"] = []
            for record in records["service_records"]:
                if not isinstance(record, dict):
                    raise BackupValidationError("Invalid service record in backup.")
                receipt_path = record.get("receipt_path")
                if receipt_path is not None:
                    if (
                        not isinstance(receipt_path, str)
                        or re.fullmatch(
                            r"receipts/receipt-[0-9]{6}\.(?:pdf|jpg|jpeg|png|webp|heic)",
                            receipt_path,
                        ) is None
                    ):
                        raise BackupValidationError("Invalid receipt reference in backup.")
                    referenced_receipts.add(receipt_path)
                receipt_filename = _backup_text(
                    record.get("receipt_filename"),
                    "receipt filename",
                    255,
                    required=False,
                )
                if (receipt_path is None) != (not receipt_filename):
                    raise BackupValidationError("Invalid receipt data in backup.")
                parsed["service_records"].append({
                    "title": _backup_text(record.get("title"), "service title", 120),
                    "category": _backup_text(record.get("category"), "service category", 24),
                    "completed_on": _backup_date(
                        record.get("completed_on"), "service date", required=True
                    ),
                    "odometer": _backup_integer(record.get("odometer"), "service mileage"),
                    "cost": _backup_decimal(record.get("cost"), "service cost"),
                    "notes": _backup_text(record.get("notes"), "service notes", 100000, required=False),
                    "receipt_filename": receipt_filename or None,
                    "receipt_path": receipt_path,
                })
            parsed["value_entries"] = []
            for entry in records["value_entries"]:
                if not isinstance(entry, dict):
                    raise BackupValidationError("Invalid value record in backup.")
                parsed["value_entries"].append({
                    "amount": _backup_decimal(entry.get("amount"), "vehicle value"),
                    "valued_on": _backup_date(entry.get("valued_on"), "value date", required=True),
                    "source": _backup_text(entry.get("source"), "value source", 120, required=False),
                    "notes": _backup_text(entry.get("notes"), "value notes", 100000, required=False),
                })
            parsed["reminders"] = []
            for reminder in records["reminders"]:
                if not isinstance(reminder, dict):
                    raise BackupValidationError("Invalid reminder in backup.")
                completed = reminder.get("completed")
                if not isinstance(completed, bool):
                    raise BackupValidationError("Invalid reminder status in backup.")
                last_notified_at = reminder.get("last_notified_at")
                if last_notified_at is not None:
                    if not isinstance(last_notified_at, str):
                        raise BackupValidationError("Invalid reminder notification date in backup.")
                    try:
                        last_notified_at = datetime.fromisoformat(last_notified_at)
                    except ValueError as error:
                        raise BackupValidationError(
                            "Invalid reminder notification date in backup."
                        ) from error
                parsed["reminders"].append({
                    "title": _backup_text(reminder.get("title"), "reminder title", 120),
                    "due_date": _backup_date(reminder.get("due_date"), "reminder date"),
                    "due_odometer": _backup_integer(reminder.get("due_odometer"), "reminder mileage"),
                    "recurrence_days": _backup_integer(reminder.get("recurrence_days"), "date recurrence"),
                    "recurrence_miles": _backup_integer(reminder.get("recurrence_miles"), "mileage recurrence"),
                    "notes": _backup_text(reminder.get("notes"), "reminder notes", 100000, required=False),
                    "completed": completed,
                    "last_notified_at": last_notified_at,
                })
            validated.append(parsed)

        if referenced_receipts != found_receipts:
            raise BackupValidationError("Receipt files in the backup do not match its records.")
        receipt_contents = {}
        for path in referenced_receipts:
            info = archive.getinfo(path)
            if info.file_size > MAX_RECEIPT_SIZE:
                raise BackupValidationError("A receipt in the backup exceeds 10 MB.")
            receipt_contents[path] = archive.read(path)
        return {
            "household_name": name,
            "discord_webhook_url": webhook_url,
            "theme": theme,
            "vehicles": validated,
            "receipt_contents": receipt_contents,
        }


def restore_backup(backup, owner_id, household, upload_folder):
    upload_path = Path(upload_folder)
    created_files = {}
    old_vehicles = Vehicle.query.filter(
        Vehicle.owner_id.in_(
            db.select(User.id).where(User.household_id == household.id)
        )
    ).all()
    old_receipts = [
        record.receipt_storage_name
        for vehicle in old_vehicles
        for record in vehicle.service_records
        if record.receipt_storage_name
    ]
    try:
        for archive_path, content in backup["receipt_contents"].items():
            suffix = Path(archive_path).suffix.lower()
            stored_name = f"{uuid.uuid4().hex}{suffix}"
            created_files[archive_path] = stored_name
            with (upload_path / stored_name).open("xb") as receipt_file:
                receipt_file.write(content)

        restored_vehicles = []
        for item in backup["vehicles"]:
            vehicle = Vehicle(
                owner_id=owner_id,
                **item["fields"],
            )
            for entry in item["odometer_entries"]:
                vehicle.odometer_entries.append(OdometerEntry(**entry))
            for record_data in item["service_records"]:
                record = dict(record_data)
                receipt_path = record.pop("receipt_path")
                record["receipt_storage_name"] = (
                    created_files.get(receipt_path) if receipt_path else None
                )
                vehicle.service_records.append(ServiceRecord(**record))
            for entry in item["value_entries"]:
                vehicle.value_entries.append(ValueEntry(**entry))
            for reminder in item["reminders"]:
                vehicle.reminders.append(Reminder(**reminder))
            restored_vehicles.append(vehicle)

        for vehicle in old_vehicles:
            db.session.delete(vehicle)
        household.name = backup["household_name"]
        household.discord_webhook_url = backup["discord_webhook_url"]
        db.session.add_all(restored_vehicles)
        current_user.theme = backup["theme"]
        db.session.commit()
    except (OSError, SQLAlchemyError):
        db.session.rollback()
        for stored_name in created_files.values():
            (upload_path / stored_name).unlink(missing_ok=True)
        raise

    for stored_name in old_receipts:
        if Path(stored_name).name == stored_name:
            try:
                (upload_path / stored_name).unlink(missing_ok=True)
            except OSError:
                logger.warning("Could not remove an old receipt after backup restore.")
    return len(restored_vehicles)


def reminder_status(reminder, current_odometer):
    if reminder.completed:
        return "Completed"
    if reminder.due_date and reminder.due_date < date.today():
        return "Overdue"
    if (
        reminder.due_odometer is not None
        and current_odometer is not None
        and reminder.due_odometer < current_odometer
    ):
        return "Overdue"
    if reminder.due_date and reminder.due_date <= date.today() + timedelta(days=30):
        return "Due soon"
    if (
        reminder.due_odometer is not None
        and current_odometer is not None
        and reminder.due_odometer <= current_odometer + 500
    ):
        return "Due soon"
    return "Upcoming"


def register_routes(app):
    @app.get("/healthz")
    def healthz():
        return {"status": "ok"}

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if current_user.is_authenticated:
            return redirect(url_for("dashboard"))
        form = LoginForm()
        if form.validate_on_submit():
            user = User.query.filter_by(username=form.username.data.strip()).first()
            if user and user.check_password(form.password.data):
                login_user(user)
                return redirect(url_for("dashboard"))
            flash("The username or password was not recognized.", "error")
        return render_template("login.html", form=form)

    @app.route("/register", methods=["GET", "POST"])
    def register():
        if current_user.is_authenticated:
            return redirect(url_for("dashboard"))

        form = AccountSetupForm()
        invitation = None
        raw_invitation = form.invite_token.data or request.args.get("invite", "")
        form.invite_token.data = raw_invitation
        if raw_invitation:
            invitation = HouseholdInvitation.query.filter_by(
                token_hash=invitation_token_hash(raw_invitation),
                accepted_at=None,
            ).filter(
                HouseholdInvitation.expires_at > utc_now()
            ).first()
            if invitation is None:
                flash("This household invitation is invalid or has expired.", "error")
                return redirect(url_for("login"))

        if form.validate_on_submit():
            user = User(
                username=form.username.data.strip(),
                household_id=invitation.household_id if invitation else None,
                role="member" if invitation else "owner",
            )
            if invitation is None:
                household = Household(name=f"{user.username}'s Garage")
                db.session.add(household)
                db.session.flush()
                user.household_id = household.id
            user.set_password(form.password.data)
            db.session.add(user)
            if invitation:
                used = db.session.execute(db.update(HouseholdInvitation).where(
                    HouseholdInvitation.id == invitation.id,
                    HouseholdInvitation.accepted_at.is_(None),
                    HouseholdInvitation.expires_at > utc_now(),
                ).values(accepted_at=utc_now()))
                if used.rowcount != 1:
                    db.session.rollback()
                    flash("This household invitation was already used or has expired.", "error")
                    return redirect(url_for("login"))
            try:
                db.session.commit()
            except IntegrityError:
                db.session.rollback()
                flash("That username is already in use. Please choose another.", "error")
                return redirect(
                    url_for("register", invite=raw_invitation) if invitation else url_for("register")
                )
            login_user(user)
            flash("Your CarFreeks account is ready.", "success")
            return redirect(url_for("dashboard"))
        household_name = invitation.household.name if invitation else None
        return render_template(
            "register.html", form=form, household_name=household_name
        )

    @app.route("/account", methods=["GET", "POST"])
    @login_required
    def account():
        form = AccountForm(username=current_user.username)
        if form.validate_on_submit():
            if not current_user.check_password(form.current_password.data):
                form.current_password.errors.append("Your current password is incorrect.")
            else:
                current_user.username = form.username.data.strip()
                if form.password.data:
                    current_user.set_password(form.password.data)
                try:
                    db.session.commit()
                except IntegrityError:
                    db.session.rollback()
                    flash("That username is already in use.", "error")
                else:
                    flash("Your account was updated.", "success")
                    return redirect(url_for("account"))
        return render_template("account.html", form=form)

    @app.post("/theme")
    @login_required
    def theme_update():
        theme = request.form.get("theme", "")
        if theme not in {"light", "dark"}:
            abort(400)
        current_user.theme = theme
        db.session.commit()
        return redirect(url_for("dashboard"))

    @app.get("/household")
    @login_required
    def household():
        members = User.query.filter_by(
            household_id=current_user.household_id
        ).order_by(User.role.desc(), User.username).all()
        invitations = HouseholdInvitation.query.filter_by(
            household_id=current_user.household_id,
            accepted_at=None,
        ).filter(
            HouseholdInvitation.expires_at > utc_now()
        ).order_by(HouseholdInvitation.created_at.desc()).all()
        new_invitation_url = session.pop("new_household_invitation_url", None)
        return render_template(
            "household.html",
            members=members,
            invitations=invitations,
            new_invitation_url=new_invitation_url,
            is_owner=is_household_owner(current_user),
            form=HouseholdForm(name=current_user.household.name),
        )

    @app.post("/household")
    @login_required
    def household_update():
        if not is_household_owner(current_user):
            abort(403)
        form = HouseholdForm()
        if form.validate_on_submit():
            webhook_url = (form.discord_webhook_url.data or "").strip()
            if not valid_discord_webhook_url(webhook_url):
                form.discord_webhook_url.errors.append(
                    "Enter a Discord webhook URL from discord.com or discordapp.com."
                )
                form.discord_webhook_url.data = ""
            else:
                current_user.household.name = form.name.data.strip()
                if form.remove_discord_webhook.data:
                    current_user.household.discord_webhook_url = None
                elif webhook_url:
                    current_user.household.discord_webhook_url = webhook_url
                db.session.commit()
                flash("Household settings saved.", "success")
                return redirect(url_for("household"))
        members = User.query.filter_by(household_id=current_user.household_id).order_by(
            User.role.desc(), User.username
        ).all()
        return render_template(
            "household.html",
            members=members,
            invitations=[],
            new_invitation_url=None,
            is_owner=True,
            form=form,
        )

    @app.post("/household/test-discord")
    @login_required
    def household_test_discord():
        if not is_household_owner(current_user):
            abort(403)
        webhook_url = current_user.household.discord_webhook_url
        if not valid_discord_webhook_url(webhook_url):
            flash("Save a valid Discord webhook URL before testing it.", "error")
        else:
            from notifications import send_discord_message

            try:
                send_discord_message(
                    webhook_url,
                    "CarFreeks Discord connection test. Your maintenance notifications are connected.",
                )
            except Exception as error:
                logger.warning("Discord test notification failed: %s", type(error).__name__)
                flash("Discord test failed. Check the webhook URL and try again.", "error")
            else:
                flash("Test message sent to Discord.", "success")
        return redirect(url_for("household"))

    @app.post("/household/invite")
    @login_required
    def household_invite():
        if not is_household_owner(current_user):
            abort(403)
        raw_token = secrets.token_urlsafe(32)
        now = utc_now()
        invitation = HouseholdInvitation(
            household_id=current_user.household_id,
            created_by_id=current_user.id,
            token_hash=invitation_token_hash(raw_token),
            created_at=now,
            expires_at=now + timedelta(days=7),
        )
        db.session.add(invitation)
        db.session.commit()
        session["new_household_invitation_url"] = url_for(
            "register", invite=raw_token, _external=True
        )
        flash("Invite link created. Share it with the person joining your household.", "success")
        return redirect(url_for("household"))

    @app.get("/exports")
    @login_required
    def exports():
        return render_template("exports.html", is_owner=is_household_owner(current_user))

    @app.get("/exports/workbook.xlsx")
    @login_required
    def export_workbook():
        return send_file(
            create_workbook(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name="carfreeks-export.xlsx",
        )

    @app.get("/exports/backup.zip")
    @login_required
    def export_backup():
        if not is_household_owner(current_user):
            abort(403)
        try:
            payload = create_backup_archive(app.config["UPLOAD_FOLDER"])
        except OSError:
            logger.exception("Could not create CarFreeks backup.")
            flash("Backup could not be created because a receipt file is unavailable.", "error")
            return redirect(url_for("exports"))
        return send_file(
            payload,
            mimetype="application/zip",
            as_attachment=True,
            download_name="carfreeks-backup.zip",
        )

    @app.post("/exports/restore")
    @login_required
    def restore_export():
        if not is_household_owner(current_user):
            abort(403)
        if request.form.get("confirm_replace") != "yes":
            flash("Confirm garage replacement before restoring a backup.", "error")
            return redirect(url_for("exports"))
        backup_file = request.files.get("backup_file")
        if not backup_file or not backup_file.filename:
            flash("Choose a CarFreeks backup ZIP file to restore.", "error")
            return redirect(url_for("exports"))
        try:
            backup = read_backup_archive(backup_file.stream)
        except (BackupValidationError, OSError, zipfile.BadZipFile) as error:
            logger.warning("Rejected CarFreeks backup: %s", error)
            flash(str(error), "error")
            return redirect(url_for("exports"))
        try:
            count = restore_backup(
                backup,
                current_user.id,
                current_user.household,
                app.config["UPLOAD_FOLDER"],
            )
        except (OSError, SQLAlchemyError):
            logger.exception("Could not restore CarFreeks backup.")
            flash("The backup could not be restored. Your existing garage was kept.", "error")
            return redirect(url_for("exports"))
        flash(f"Backup restored. {count} vehicle(s) are now in your garage.", "success")
        return redirect(url_for("dashboard"))

    @app.post("/household/invitations/<int:invitation_id>/revoke")
    @login_required
    def household_invitation_revoke(invitation_id):
        if not is_household_owner(current_user):
            abort(403)
        invitation = HouseholdInvitation.query.filter_by(
            id=invitation_id,
            household_id=current_user.household_id,
            accepted_at=None,
        ).first_or_404()
        db.session.delete(invitation)
        db.session.commit()
        flash("Household invite revoked.", "success")
        return redirect(url_for("household"))

    @app.post("/logout")
    @login_required
    def logout():
        logout_user()
        flash("You have signed out.", "info")
        return redirect(url_for("login"))

    @app.get("/")
    @login_required
    def dashboard():
        member_ids = db.select(User.id).where(
            User.household_id == current_user.household_id
        )
        vehicles = Vehicle.query.filter(Vehicle.owner_id.in_(member_ids)).order_by(
            Vehicle.nickname
        ).all()
        reminders = []
        for vehicle in vehicles:
            for reminder in vehicle.reminders:
                if not reminder.completed:
                    reminders.append(
                        (vehicle, reminder, reminder_status(reminder, vehicle.current_odometer))
                    )
        reminders.sort(
            key=lambda item: (
                item[1].due_date or date.max,
                item[1].due_odometer or 2**63,
                item[0].nickname.lower(),
            )
        )
        return render_template(
            "dashboard.html", vehicles=vehicles, reminders=reminders[:8]
        )

    @app.route("/vehicles/new", methods=["GET", "POST"])
    @login_required
    def vehicle_new():
        form = VehicleForm()
        if form.validate_on_submit():
            vehicle = Vehicle(owner_id=current_user.id)
            form.populate_obj(vehicle)
            db.session.add(vehicle)
            db.session.commit()
            flash(f"{vehicle.nickname} was added.", "success")
            return redirect(url_for("vehicle_detail", vehicle_id=vehicle.id))
        return render_template("vehicle_form.html", form=form, heading="Add a vehicle")

    @app.route("/vehicles/<int:vehicle_id>/edit", methods=["GET", "POST"])
    @login_required
    def vehicle_edit(vehicle_id):
        vehicle = get_vehicle_or_404(vehicle_id)
        form = VehicleForm(obj=vehicle)
        if form.validate_on_submit():
            form.populate_obj(vehicle)
            db.session.commit()
            flash(f"{vehicle.nickname} was updated.", "success")
            return redirect(url_for("vehicle_detail", vehicle_id=vehicle.id))
        return render_template("vehicle_form.html", form=form, heading="Edit vehicle")

    @app.post("/vehicles/<int:vehicle_id>/delete")
    @login_required
    def vehicle_delete(vehicle_id):
        vehicle = get_vehicle_or_404(vehicle_id)
        receipt_names = [
            record.receipt_storage_name
            for record in vehicle.service_records
            if record.receipt_storage_name
        ]
        name = vehicle.nickname
        db.session.delete(vehicle)
        db.session.commit()
        for stored_name in receipt_names:
            (Path(app.config["UPLOAD_FOLDER"]) / stored_name).unlink(missing_ok=True)
        flash(f"{name} and its records were deleted.", "success")
        return redirect(url_for("dashboard"))

    @app.get("/vehicles/<int:vehicle_id>")
    @login_required
    def vehicle_detail(vehicle_id):
        vehicle = get_vehicle_or_404(vehicle_id)
        records = sorted(
            vehicle.service_records,
            key=lambda item: (item.completed_on, item.id),
            reverse=True,
        )
        values = sorted(
            vehicle.value_entries,
            key=lambda item: (item.valued_on, item.id),
            reverse=True,
        )
        odometer_entries = sorted(
            vehicle.odometer_entries,
            key=lambda item: (item.recorded_on, item.id),
            reverse=True,
        )
        reminders = sorted(
            vehicle.reminders,
            key=lambda item: (
                item.completed,
                item.due_date or date.max,
                item.due_odometer or 2**63,
                item.id,
            ),
        )
        return render_template(
            "vehicle_detail.html",
            vehicle=vehicle,
            records=records,
            values=values,
            odometer_entries=odometer_entries,
            reminders=reminders,
            reminder_status=reminder_status,
        )

    @app.route("/vehicles/<int:vehicle_id>/odometer/new", methods=["GET", "POST"])
    @login_required
    def odometer_new(vehicle_id):
        vehicle = get_vehicle_or_404(vehicle_id)
        form = OdometerForm()
        if form.validate_on_submit():
            entry = OdometerEntry(vehicle_id=vehicle.id)
            form.populate_obj(entry)
            db.session.add(entry)
            db.session.commit()
            flash("Odometer reading saved.", "success")
            return redirect(url_for("vehicle_detail", vehicle_id=vehicle.id))
        return render_template(
            "simple_form.html",
            form=form,
            heading=f"Add mileage · {vehicle.nickname}",
            back_url=url_for("vehicle_detail", vehicle_id=vehicle.id),
        )

    @app.route("/vehicles/<int:vehicle_id>/service/new", methods=["GET", "POST"])
    @login_required
    def service_new(vehicle_id):
        vehicle = get_vehicle_or_404(vehicle_id)
        form = ServiceForm()
        if form.validate_on_submit():
            receipt = form.receipt.data
            if receipt:
                receipt.stream.seek(0, os.SEEK_END)
                receipt_size = receipt.stream.tell()
                receipt.stream.seek(0)
                if receipt_size > MAX_RECEIPT_SIZE:
                    flash("Receipt files must be 10 MB or smaller.", "error")
                    return render_template(
                        "simple_form.html",
                        form=form,
                        heading=f"Add service or repair · {vehicle.nickname}",
                        back_url=url_for("vehicle_detail", vehicle_id=vehicle.id),
                    )
            original_name = secure_filename(receipt.filename) if receipt else ""
            if receipt and original_name:
                extension = Path(original_name).suffix.lower()
                if extension not in ALLOWED_RECEIPT_EXTENSIONS:
                    flash("Use a PDF or supported image file for a receipt.", "error")
                    return render_template(
                        "simple_form.html",
                        form=form,
                        heading=f"Add service or repair · {vehicle.nickname}",
                        back_url=url_for("vehicle_detail", vehicle_id=vehicle.id),
                    )
                stored_name = f"{uuid.uuid4().hex}{extension}"
                receipt.save(Path(app.config["UPLOAD_FOLDER"]) / stored_name)
            else:
                original_name = None
                stored_name = None

            record = ServiceRecord(vehicle_id=vehicle.id)
            form.populate_obj(record)
            record.receipt_filename = original_name
            record.receipt_storage_name = stored_name
            db.session.add(record)
            db.session.commit()
            flash("Service or repair record saved.", "success")
            return redirect(url_for("vehicle_detail", vehicle_id=vehicle.id))
        return render_template(
            "simple_form.html",
            form=form,
            heading=f"Add service or repair · {vehicle.nickname}",
            back_url=url_for("vehicle_detail", vehicle_id=vehicle.id),
        )

    @app.get("/receipts/<int:record_id>")
    @login_required
    def receipt_download(record_id):
        record = (
            ServiceRecord.query.join(Vehicle)
            .filter(
                ServiceRecord.id == record_id,
                Vehicle.owner_id.in_(
                    db.select(User.id).where(
                        User.household_id == current_user.household_id
                    )
                ),
            )
            .first_or_404()
        )
        if not record.receipt_storage_name:
            abort(404)
        return send_from_directory(
            app.config["UPLOAD_FOLDER"],
            record.receipt_storage_name,
            as_attachment=True,
            download_name=record.receipt_filename or "receipt",
        )

    @app.route("/vehicles/<int:vehicle_id>/value/new", methods=["GET", "POST"])
    @login_required
    def value_new(vehicle_id):
        vehicle = get_vehicle_or_404(vehicle_id)
        form = ValueForm()
        if form.validate_on_submit():
            entry = ValueEntry(vehicle_id=vehicle.id)
            form.populate_obj(entry)
            db.session.add(entry)
            db.session.commit()
            flash("Vehicle value saved to history.", "success")
            return redirect(url_for("vehicle_detail", vehicle_id=vehicle.id))
        return render_template(
            "simple_form.html",
            form=form,
            heading=f"Record estimated value · {vehicle.nickname}",
            back_url=url_for("vehicle_detail", vehicle_id=vehicle.id),
        )

    @app.route("/vehicles/<int:vehicle_id>/reminders/new", methods=["GET", "POST"])
    @login_required
    def reminder_new(vehicle_id):
        vehicle = get_vehicle_or_404(vehicle_id)
        form = ReminderForm()
        if form.validate_on_submit():
            reminder = Reminder(vehicle_id=vehicle.id)
            form.populate_obj(reminder)
            db.session.add(reminder)
            db.session.commit()
            flash("Maintenance reminder saved.", "success")
            return redirect(url_for("vehicle_detail", vehicle_id=vehicle.id))
        return render_template(
            "simple_form.html",
            form=form,
            heading=f"Add maintenance reminder · {vehicle.nickname}",
            back_url=url_for("vehicle_detail", vehicle_id=vehicle.id),
        )

    @app.post("/vehicles/<int:vehicle_id>/reminders/<int:reminder_id>/toggle")
    @login_required
    def reminder_toggle(vehicle_id, reminder_id):
        vehicle = get_vehicle_or_404(vehicle_id)
        reminder = Reminder.query.filter_by(
            id=reminder_id, vehicle_id=vehicle.id
        ).first_or_404()
        if reminder.completed:
            reminder.completed = False
            db.session.commit()
            flash("Reminder reopened.", "success")
        elif reminder.recurrence_days or reminder.recurrence_miles:
            today = date.today()
            if reminder.recurrence_days:
                next_date = reminder.due_date or today
                while next_date <= today:
                    next_date += timedelta(days=reminder.recurrence_days)
                reminder.due_date = next_date
            elif reminder.recurrence_miles:
                reminder.due_date = None
            if reminder.recurrence_miles:
                current_odometer = vehicle.current_odometer or 0
                next_mileage = reminder.due_odometer or current_odometer
                while next_mileage <= current_odometer:
                    next_mileage += reminder.recurrence_miles
                reminder.due_odometer = next_mileage
            elif reminder.recurrence_days:
                reminder.due_odometer = None
            reminder.last_notified_at = None
            db.session.commit()
            flash("Maintenance marked done; your next recurring service is scheduled.", "success")
        else:
            reminder.completed = True
            db.session.commit()
            flash("Reminder marked complete.", "success")
        return redirect(url_for("vehicle_detail", vehicle_id=vehicle.id))

    @app.errorhandler(413)
    def upload_too_large(_error):
        flash("This upload is too large. Receipts are limited to 10 MB and backups to 100 MB.", "error")
        return redirect(url_for("exports"))


if __name__ == "__main__":
    create_app().run(host="0.0.0.0", port=8080)
